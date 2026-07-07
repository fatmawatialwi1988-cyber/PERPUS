import customtkinter as ctk
from database.models import BookModel
from database.db import get_connection
from barcode.generator import BarcodeGenerator
import openpyxl
from tkinter import filedialog, messagebox
import os

class BooksFrame(ctk.CTkFrame):
    def __init__(self, parent, controller):
        super().__init__(parent)
        self.controller = controller

        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(2, weight=1)

        # Header Title and Toolbar
        self.header_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.header_frame.grid(row=0, column=0, padx=20, pady=15, sticky="ew")
        
        self.header_label = ctk.CTkLabel(self.header_frame, text="KELOLA DATA BUKU", font=ctk.CTkFont(size=20, weight="bold"))
        self.header_label.pack(side="left")

        # Top search and filter bar
        self.toolbar = ctk.CTkFrame(self, fg_color="transparent")
        self.toolbar.grid(row=1, column=0, padx=20, pady=5, sticky="ew")

        self.search_entry = ctk.CTkEntry(self.toolbar, placeholder_text="Cari judul, pengarang, ISBN...")
        self.search_entry.pack(side="left", padx=5)
        self.search_entry.bind("<KeyRelease>", lambda e: self.load_books())

        # Category and Rack dropdown filters
        self.cat_filter = ctk.CTkComboBox(self.toolbar, values=["Semua Kategori"], command=lambda e: self.load_books())
        self.cat_filter.pack(side="left", padx=5)

        self.rack_filter = ctk.CTkComboBox(self.toolbar, values=["Semua Rak"], command=lambda e: self.load_books())
        self.rack_filter.pack(side="left", padx=5)

        # Right-side buttons
        self.btn_add = ctk.CTkButton(self.toolbar, text="+ Tambah Buku", command=self.open_add_dialog, width=100)
        self.btn_add.pack(side="right", padx=5)

        self.btn_export = ctk.CTkButton(self.toolbar, text="Ekspor Excel", fg_color="#059669", hover_color="#047857", command=self.export_excel, width=100)
        self.btn_export.pack(side="right", padx=5)

        self.btn_import = ctk.CTkButton(self.toolbar, text="Impor Excel", fg_color="#10b981", hover_color="#059669", command=self.import_excel, width=100)
        self.btn_import.pack(side="right", padx=5)

        # 3. Main Data Container (Books Table / List View)
        self.list_frame = ctk.CTkScrollableFrame(self)
        self.list_frame.grid(row=2, column=0, padx=20, pady=10, sticky="nsew")

        # Setup lists
        self.populate_filters()
        self.load_books()

    def populate_filters(self):
        conn = get_connection()
        cursor = conn.cursor()
        
        cursor.execute("SELECT id, nama FROM kategori;")
        self.categories = cursor.fetchall()
        cat_vals = ["Semua Kategori"] + [c['nama'] for c in self.categories]
        self.cat_filter.configure(values=cat_vals)
        self.cat_filter.set("Semua Kategori")

        cursor.execute("SELECT id, nama FROM rak;")
        self.racks = cursor.fetchall()
        rack_vals = ["Semua Rak"] + [r['nama'] for r in self.racks]
        self.rack_filter.configure(values=rack_vals)
        self.rack_filter.set("Semua Rak")
        conn.close()

    def load_books(self):
        # Clear list
        for child in self.list_frame.winfo_children():
            child.destroy()

        # Get active filters
        search = self.search_entry.get()
        
        selected_cat_name = self.cat_filter.get()
        cat_id = None
        if selected_cat_name != "Semua Kategori":
            match = [c['id'] for c in self.categories if c['nama'] == selected_cat_name]
            if match: cat_id = match[0]

        selected_rack_name = self.rack_filter.get()
        rack_id = None
        if selected_rack_name != "Semua Rak":
            match = [r['id'] for r in self.racks if r['nama'] == selected_rack_name]
            if match: rack_id = match[0]

        books = BookModel.get_all(search, cat_id, rack_id)

        # Header Row
        headers = ["Kode", "Judul", "Pengarang", "Penerbit", "Tahun", "Stok", "Tersedia", "Aksi"]
        for col_idx, h in enumerate(headers):
            lbl = ctk.CTkLabel(self.list_frame, text=h, font=ctk.CTkFont(weight="bold", size=12))
            lbl.grid(row=0, column=col_idx, padx=10, pady=5, sticky="w")

        for row_idx, b in enumerate(books):
            ctk.CTkLabel(self.list_frame, text=b['kode_buku'], font=ctk.CTkFont(family="monospace")).grid(row=row_idx+1, column=0, padx=10, pady=3, sticky="w")
            ctk.CTkLabel(self.list_frame, text=b['judul'][:24], font=ctk.CTkFont(weight="bold")).grid(row=row_idx+1, column=1, padx=10, pady=3, sticky="w")
            ctk.CTkLabel(self.list_frame, text=b['pengarang'][:15]).grid(row=row_idx+1, column=2, padx=10, pady=3, sticky="w")
            ctk.CTkLabel(self.list_frame, text=b['penerbit'][:15]).grid(row=row_idx+1, column=3, padx=10, pady=3, sticky="w")
            ctk.CTkLabel(self.list_frame, text=str(b['tahun'])).grid(row=row_idx+1, column=4, padx=10, pady=3, sticky="w")
            ctk.CTkLabel(self.list_frame, text=f"{b['jumlah_buku']} Unit").grid(row=row_idx+1, column=5, padx=10, pady=3, sticky="w")
            ctk.CTkLabel(self.list_frame, text=f"{b['jumlah_tersedia']} Unit", text_color="#10b981" if b['jumlah_tersedia'] > 0 else "#dc2626").grid(row=row_idx+1, column=6, padx=10, pady=3, sticky="w")

            # Actions cell
            act_box = ctk.CTkFrame(self.list_frame, fg_color="transparent")
            act_box.grid(row=row_idx+1, column=7, padx=10, pady=3, sticky="e")

            ctk.CTkButton(act_box, text="Edit", width=40, command=lambda bk=b: self.open_edit_dialog(bk)).pack(side="left", padx=2)
            ctk.CTkButton(act_box, text="Hapus", width=40, fg_color="#dc2626", hover_color="#991b1b", command=lambda k=b['kode_buku']: self.delete_book(k)).pack(side="left", padx=2)
            ctk.CTkButton(act_box, text="Barcode", width=45, fg_color="#4f46e5", hover_color="#4338ca", command=lambda k=b['kode_buku'], bcd=b['barcode']: self.print_barcode(k, bcd)).pack(side="left", padx=2)

    def delete_book(self, code_book):
        if self.controller.active_user['role'] != 'Administrator':
            messagebox.showerror("Akses Ditolak", "Hanya Administrator yang diperbolehkan menghapus data buku!")
            return
        if messagebox.askyesno("Hapus Buku", f"Apakah Anda yakin menghapus buku {code_book}?"):
            if BookModel.delete(code_book):
                messagebox.showinfo("Sukses", "Buku berhasil dihapus dari katalog.")
                self.load_books()
            else:
                messagebox.showerror("Error", "Buku tidak dapat dihapus (mungkin masih dalam status dipinjam).")

    def print_barcode(self, kode, barcode):
        # Fetch book details from SQLite
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT judul, isbn FROM buku WHERE kode_buku = ?;", (kode,))
        row = cursor.fetchone()
        conn.close()
        
        judul = row['judul'] if row else "Katalog Buku"
        isbn = row['isbn'] if row else "-"
        
        os.makedirs("barcodes", exist_ok=True)
        os.makedirs("barcodes/books", exist_ok=True)
        
        filepath = f"barcodes/books/barcode_{kode}.jpg"
        
        try:
            BarcodeGenerator.generate_book_barcode_jpg(kode, judul, isbn, filepath)
            messagebox.showinfo("Download Barcode JPG", f"Barcode Cetak JPG (300 DPI) untuk buku {kode} berhasil dibuat!\n\nFile disimpan pada folder:\n{filepath}")
        except Exception as e:
            messagebox.showerror("Error", f"Gagal membuat Barcode Cetak: {e}")

    def open_add_dialog(self):
        self.open_book_form(None)

    def open_edit_dialog(self, book):
        self.open_book_form(book)

    def open_book_form(self, book_data=None):
        # Create small dialog window
        dialog = ctk.CTkToplevel(self)
        dialog.title("Form Input Buku")
        dialog.geometry("400x550")
        dialog.transient(self)
        dialog.grab_set()

        ctk.CTkLabel(dialog, text="FORM DATA BUKU", font=ctk.CTkFont(size=16, weight="bold")).pack(pady=15)

        frame = ctk.CTkFrame(dialog, fg_color="transparent")
        frame.pack(fill="both", expand=True, padx=20, pady=5)

        # Fields
        lbl_kode = ctk.CTkLabel(frame, text="Kode Buku (Unik) *")
        lbl_kode.pack(anchor="w")
        ent_kode = ctk.CTkEntry(frame, placeholder_text="B001")
        ent_kode.pack(fill="x", pady=(0, 10))
        if book_data:
            ent_kode.insert(0, book_data['kode_buku'])
            ent_kode.configure(state="disabled")

        lbl_isbn = ctk.CTkLabel(frame, text="ISBN / Barcode ID *")
        lbl_isbn.pack(anchor="w")
        ent_isbn = ctk.CTkEntry(frame, placeholder_text="9789791234561")
        ent_isbn.pack(fill="x", pady=(0, 10))
        if book_data:
            ent_isbn.insert(0, book_data['isbn'])

        lbl_judul = ctk.CTkLabel(frame, text="Judul Buku *")
        lbl_judul.pack(anchor="w")
        ent_judul = ctk.CTkEntry(frame, placeholder_text="Laskar Pelangi")
        ent_judul.pack(fill="x", pady=(0, 10))
        if book_data:
            ent_judul.insert(0, book_data['judul'])

        lbl_author = ctk.CTkLabel(frame, text="Pengarang *")
        lbl_author.pack(anchor="w")
        ent_author = ctk.CTkEntry(frame, placeholder_text="Andrea Hirata")
        ent_author.pack(fill="x", pady=(0, 10))
        if book_data:
            ent_author.insert(0, book_data['pengarang'])

        lbl_publisher = ctk.CTkLabel(frame, text="Penerbit *")
        lbl_publisher.pack(anchor="w")
        ent_publisher = ctk.CTkEntry(frame, placeholder_text="Bentang Pustaka")
        ent_publisher.pack(fill="x", pady=(0, 10))
        if book_data:
            ent_publisher.insert(0, book_data['penerbit'])

        # Multi-columns
        sub_frame = ctk.CTkFrame(frame, fg_color="transparent")
        sub_frame.pack(fill="x", pady=(0, 10))
        sub_frame.grid_columnconfigure(0, weight=1)
        sub_frame.grid_columnconfigure(1, weight=1)

        lbl_year = ctk.CTkLabel(sub_frame, text="Tahun")
        lbl_year.grid(row=0, column=0, sticky="w")
        ent_year = ctk.CTkEntry(sub_frame, placeholder_text="2026")
        ent_year.grid(row=1, column=0, sticky="ew", padx=(0, 5))
        if book_data:
            ent_year.insert(0, str(book_data['tahun']))
        else:
            ent_year.insert(0, "2026")

        lbl_stock = ctk.CTkLabel(sub_frame, text="Jumlah Stok")
        lbl_stock.grid(row=0, column=1, sticky="w")
        ent_stock = ctk.CTkEntry(sub_frame, placeholder_text="5")
        ent_stock.grid(row=1, column=1, sticky="ew", padx=(5, 0))
        if book_data:
            ent_stock.insert(0, str(book_data['jumlah_buku']))
        else:
            ent_stock.insert(0, "5")

        # Save trigger
        def save():
            kode = ent_kode.get().strip().upper()
            isbn = ent_isbn.get().strip()
            judul = ent_judul.get().strip()
            author = ent_author.get().strip()
            pub = ent_publisher.get().strip()
            
            try:
                tahun = int(ent_year.get().strip() or 2026)
                stock = int(ent_stock.get().strip() or 5)
            except ValueError:
                messagebox.showerror("Error", "Tahun dan stok harus diisi berupa angka!")
                return

            if not kode or not isbn or not judul or not author or not pub:
                messagebox.showerror("Error", "Isi seluruh data berlabel wajib (*)")
                return

            # categories / rack placeholders
            cat_id = self.categories[0]['id'] if self.categories else 1
            rack_id = self.racks[0]['id'] if self.racks else 1

            if book_data:
                # Update
                success = BookModel.update(kode, isbn, kode, judul, author, pub, tahun, cat_id, rack_id, stock)
            else:
                # Insert
                success = BookModel.add(kode, isbn, kode, judul, author, pub, tahun, cat_id, rack_id, stock)

            if success:
                messagebox.showinfo("Sukses", "Data buku berhasil disimpan!")
                self.load_books()
                dialog.destroy()
            else:
                messagebox.showerror("Error", "Gagal menyimpan data buku. Kode Buku atau ISBN mungkin duplikat.")

        btn_save = ctk.CTkButton(dialog, text="Simpan Buku", command=save)
        btn_save.pack(pady=15)

    def export_excel(self):
        filepath = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel Files", "*.xlsx")])
        if not filepath:
            return

        # build workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Katalog Buku"

        # Headers
        headers = ["Kode Buku", "ISBN", "Barcode ID", "Judul Buku", "Pengarang", "Penerbit", "Tahun", "Jumlah Buku", "Tersedia", "Dipinjam"]
        ws.append(headers)

        books = BookModel.get_all()
        for b in books:
            ws.append([
                b['kode_buku'], b['isbn'], b['barcode'], b['judul'], b['pengarang'],
                b['penerbit'], b['tahun'], b['jumlah_buku'], b['jumlah_tersedia'], b['jumlah_dipinjam']
            ])

        wb.save(filepath)
        messagebox.showinfo("Ekspor Excel", f"Data buku berhasil diekspor ke {filepath}")

    def import_excel(self):
        filepath = filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx")])
        if not filepath:
            return

        try:
            wb = openpyxl.load_workbook(filepath)
            ws = wb.active
            rows = list(ws.rows)
            if len(rows) <= 1:
                messagebox.showerror("Error", "Berkas Excel kosong!")
                return

            conn = get_connection()
            cursor = conn.cursor()
            imported = 0

            # first row is header
            for r in rows[1:]:
                vals = [cell.value for cell in r]
                if len(vals) < 5 or not vals[0]: continue
                
                # Check duplicate
                cursor.execute("SELECT COUNT(*) FROM buku WHERE kode_buku = ?;", (str(vals[0]).strip().upper(),))
                if cursor.fetchone()[0] > 0: continue

                # placeholder rack/category
                cat_id = self.categories[0]['id'] if self.categories else 1
                rack_id = self.racks[0]['id'] if self.racks else 1

                cursor.execute("""
                INSERT INTO buku (kode_buku, isbn, barcode, judul, pengarang, penerbit, tahun, kategori_id, rak_id, jumlah_buku, jumlah_tersedia, jumlah_dipinjam, status)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, 'Tersedia');
                """, (
                    str(vals[0]).strip().upper(),
                    str(vals[1] or vals[0]).strip(),
                    str(vals[2] or vals[0]).strip(),
                    str(vals[3] or 'Judul').strip(),
                    str(vals[4] or 'Pengarang').strip(),
                    str(vals[5] or 'Penerbit').strip(),
                    int(vals[6] or 2026),
                    cat_id,
                    rack_id,
                    int(vals[7] or 5),
                    int(vals[7] or 5)
                ))
                imported += 1

            conn.commit()
            conn.close()
            messagebox.showinfo("Impor Berhasil", f"Berhasil mengimpor {imported} buku baru ke database.")
            self.load_books()
        except Exception as e:
            messagebox.showerror("Error Impor", f"Gagal membaca Excel: {e}")
